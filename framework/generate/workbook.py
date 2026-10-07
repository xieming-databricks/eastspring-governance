"""Read an Access-Matrix Excel workbook into plain row-dicts (openpyxl I/O only).

``read_workbook(path)`` returns ``{tab_name: [ {header: value}, ... ]}`` for the requested
tabs. This is the only module that touches openpyxl / the filesystem; the transform in
``generator`` is pure and consumes these dicts, so it is testable without Excel.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

import openpyxl


def _clean(value: Any) -> Any:
    """Strip surrounding whitespace from strings; pass other values through."""
    return value.strip() if isinstance(value, str) else value


def _read_sheet(ws) -> list[dict[str, Any]]:
    """First non-blank row is the header; return the remaining non-blank rows as dicts."""
    headers: list[str] | None = None
    rows: list[dict[str, Any]] = []
    for raw in ws.iter_rows(values_only=True):
        if not any(c is not None and str(c).strip() for c in raw):
            continue                                    # skip fully blank rows
        if headers is None:
            headers = [_clean(c) for c in raw]
            continue
        rec = {h: _clean(v) for h, v in zip(headers, raw) if h}   # drop unnamed columns
        rows.append(rec)
    return rows


def read_workbook(path: str | Path, sheets: list[str] | tuple[str, ...] | None = None
                  ) -> dict[str, list[dict[str, Any]]]:
    """Load ``sheets`` (default: all) from the workbook at ``path`` into row-dicts.

    A requested sheet that is absent is simply omitted (the generator reports the gap)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        names = list(sheets) if sheets is not None else list(wb.sheetnames)
        return {name: _read_sheet(wb[name]) for name in names if name in wb.sheetnames}
    finally:
        wb.close()
